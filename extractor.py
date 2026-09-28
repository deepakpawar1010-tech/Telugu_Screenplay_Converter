# -*- coding: utf-8 -*-
"""
Document Extractor for Telugu Screenplay Converter
Extracts text and screenplay blocks from DOCX and PDF screenplay files.
"""

import os
import io
from typing import List, Union
import docx
import pymupdf

from screenplay_parser import parse_screenplay_text, ScreenplayBlock


def extract_text_from_docx(source: Union[str, bytes, io.BytesIO]) -> str:
    """
    Extracts text from a DOCX file path or bytes buffer.
    Preserves screenplay paragraph breaks.
    """
    if isinstance(source, bytes):
        doc = docx.Document(io.BytesIO(source))
    elif isinstance(source, io.BytesIO):
        doc = docx.Document(source)
    else:
        doc = docx.Document(source)

    paragraphs = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            paragraphs.append(txt)

    return "\n\n".join(paragraphs)


def extract_text_from_pdf(source: Union[str, bytes, io.BytesIO]) -> str:
    """
    Extracts text from a PDF file path or bytes buffer using PyMuPDF.
    Cleans up common header/footer line artifacts.
    """
    if isinstance(source, bytes):
        doc = pymupdf.open(stream=source, filetype="pdf")
    elif isinstance(source, io.BytesIO):
        doc = pymupdf.open(stream=source.getvalue(), filetype="pdf")
    else:
        doc = pymupdf.open(source)

    extracted_pages = []
    for page in doc:
        text = page.get_text("text")
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        
        # Filter out standalone page number lines like "Page 12" or "12"
        filtered_lines = []
        for line in lines:
            if line.isdigit() and len(line) <= 3:
                continue
            if line.lower().startswith("page ") and line[5:].strip().isdigit():
                continue
            filtered_lines.append(line)

        if filtered_lines:
            extracted_pages.append("\n\n".join(filtered_lines))

    doc.close()
    return "\n\n".join(extracted_pages)


def extract_screenplay_blocks(
    source: Union[str, bytes, io.BytesIO], 
    file_type: str = "docx"
) -> List[ScreenplayBlock]:
    """
    Extracts and parses screenplay blocks directly from a file or buffer.
    file_type can be 'docx', 'pdf', or 'txt'.
    """
    ft = file_type.lower().lstrip(".")
    if ft == "docx":
        text = extract_text_from_docx(source)
    elif ft == "pdf":
        text = extract_text_from_pdf(source)
    elif ft in ("txt", "text"):
        if isinstance(source, bytes):
            text = source.decode("utf-8", errors="replace")
        elif isinstance(source, io.BytesIO):
            text = source.getvalue().decode("utf-8", errors="replace")
        elif os.path.exists(source):
            with open(source, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()
        else:
            text = str(source)
    else:
        raise ValueError(f"Unsupported file type: {file_type}. Supported: docx, pdf, txt")

    return parse_screenplay_text(text)
